import os
import tempfile
import pytest
from src.db.connection import get_db_connection, immediate_transaction
from src.db.init_db import init_db
from src.services.stock_service import add_product_service
from src.services.billing_service import start_bill_service, add_item_to_bill_service, finalize_bill_service, get_bill_details_service
from src.documents.pdf_invoice import generate_pdf_invoice_file
from src.documents.pptx_deck import generate_pptx_analysis_deck_file

@pytest.fixture
def test_db():
    with tempfile.TemporaryDirectory() as tmpdir:
        db_path = os.path.join(tmpdir, "test_docs.db")
        init_db(db_path)
        yield db_path

def test_pdf_invoice_generation(test_db):
    from src.documents.pdf_invoice import get_invoice_font_names
    norm_font, bold_font = get_invoice_font_names()
    assert norm_font != "Helvetica", "A Unicode TrueType font must be registered for PDF currency rendering"

    with get_db_connection(test_db) as conn:
        with immediate_transaction(conn):
            add_product_service(conn, {
                "name": "Tea 250g", "brand": "BrandT", "unit": "packet", "is_loose": False,
                "hsn_code": "0902", "gst_slab": 5, "cost_price": 80, "sell_price": 100,
                "mrp": 110, "quantity": 50, "reorder_level": 5
            })

        with immediate_transaction(conn):
            bill = start_bill_service(conn, "Anil Kumar", payment_mode="upi")
            add_item_to_bill_service(conn, bill["bill"]["id"], "Tea 250g", quantity=2)
            fin_bill = finalize_bill_service(conn, bill["bill"]["id"], payment_reference="UPI998877")

        bill_details = get_bill_details_service(conn, fin_bill["bill"]["id"])
        prefs = {"shop_name": "Testing Kirana", "shop_gstin": "33AAAAA1111B1Z2"}

        pdf_path = generate_pdf_invoice_file(bill_details, prefs)
        assert os.path.exists(pdf_path)
        assert os.path.getsize(pdf_path) > 1000

def test_pdf_invoice_rupee_font_embedding(test_db):
    import re
    from src.documents.pdf_invoice import generate_pdf_invoice_file, get_invoice_font_names

    norm_font, bold_font = get_invoice_font_names()
    assert "Helvetica" not in norm_font, "Helvetica must not be used as fallback"

    with get_db_connection(test_db) as conn:
        with immediate_transaction(conn):
            add_product_service(conn, {
                "name": "Test Item 1", "brand": "Brand A", "unit": "packet", "is_loose": False,
                "hsn_code": "1001", "gst_slab": 12, "cost_price": 10, "sell_price": 14,
                "mrp": 15, "quantity": 100, "reorder_level": 5
            })
        with immediate_transaction(conn):
            b = start_bill_service(conn, "Rohan Sharma", payment_mode="upi")
            add_item_to_bill_service(conn, b["bill"]["id"], "Test Item 1", quantity=6)
            fin = finalize_bill_service(conn, b["bill"]["id"], payment_reference="UPI9876543210")

        bill_details = get_bill_details_service(conn, fin["bill"]["id"])
        prefs = {"shop_name": "Testing Kirana"}
        pdf_path = generate_pdf_invoice_file(bill_details, prefs)

    assert os.path.exists(pdf_path)

    with open(pdf_path, 'rb') as f:
        content = f.read().decode('latin-1')

    base_fonts = set(re.findall(r'/BaseFont\s*/([A-Za-z0-9\+\-]+)', content))

    for forbidden in ['Helvetica', 'Helvetica-Bold', 'Helvetica-Oblique', 'ZapfDingbats', 'Symbol']:
        assert not any(forbidden in font_name for font_name in base_fonts), f"Forbidden font '{forbidden}' embedded in PDF: {base_fonts}"

    assert any("DejaVu" in f or "FreeSans" in f or "Segoe" in f or "Arial" in f for f in base_fonts), f"Unicode TTF font missing in PDF: {base_fonts}"

def test_pptx_deck_generation_with_data_and_empty_range(test_db):
    from pptx import Presentation
    with get_db_connection(test_db) as conn:
        # 1. Empty Date Range Test
        pptx_empty = generate_pptx_analysis_deck_file(conn, "2020-01-01", "2020-01-02")
        assert os.path.exists(pptx_empty)
        assert os.path.getsize(pptx_empty) > 1000
        prs_empty = Presentation(pptx_empty)
        assert len(prs_empty.slides) == 1

        # 2. Date Range with Data
        with immediate_transaction(conn):
            add_product_service(conn, {
                "name": "Milk 1L", "brand": "BrandM", "unit": "litre", "is_loose": False,
                "hsn_code": "0401", "gst_slab": 0, "cost_price": 45, "sell_price": 54,
                "mrp": 55, "quantity": 100, "reorder_level": 10
            })
            b = start_bill_service(conn, "Sunil")
            add_item_to_bill_service(conn, b["bill"]["id"], "Milk 1L", 3)
            finalize_bill_service(conn, b["bill"]["id"])

        import datetime
        today = str(datetime.date.today())
        pptx_data = generate_pptx_analysis_deck_file(conn, today, today)
        assert os.path.exists(pptx_data)
        assert os.path.getsize(pptx_data) > 1000

        prs_data = Presentation(pptx_data)
        assert len(prs_data.slides) == 5
        slide5_text = prs_data.slides[4].shapes[0].text_frame.text
        assert "GST Breakdown" in slide5_text

def test_generate_invoice_pdf_gemini_supplied_unmentioned_bill_id_ignored(test_db):
    from src.agent.context import current_chat_id_var, current_user_message_var
    from src.tools.report_tools import generate_invoice_pdf

    current_chat_id_var.set("chat_test_unmentioned")
    current_user_message_var.set("Generate the PDF invoice for my latest bill")

    with get_db_connection(test_db) as conn:
        with immediate_transaction(conn):
            add_product_service(conn, {
                "name": "Item Unmentioned", "brand": "BrandU", "unit": "packet", "is_loose": False,
                "hsn_code": "1111", "gst_slab": 5, "cost_price": 50, "sell_price": 100,
                "mrp": 100, "quantity": 10, "reorder_level": 2
            })
            b1 = start_bill_service(conn, "Customer 1", chat_id="chat_test_unmentioned")
            add_item_to_bill_service(conn, b1["bill"]["id"], "Item Unmentioned", quantity=1)
            fin1 = finalize_bill_service(conn, b1["bill"]["id"])

            b2 = start_bill_service(conn, "Customer 2", chat_id="chat_test_unmentioned")
            add_item_to_bill_service(conn, b2["bill"]["id"], "Item Unmentioned", quantity=2)
            fin2 = finalize_bill_service(conn, b2["bill"]["id"])

    # User prompt said "Generate the PDF invoice for my latest bill" (did NOT mention bill b1)
    # Gemini sends bill_id=fin1["bill"]["id"]
    res = generate_invoice_pdf(bill_id=fin1["bill"]["id"], db_path=test_db)
    # Resolver MUST ignore fin1 and return fin2 (the latest bill for current chat)
    assert res["status"] == "success"
    assert res["bill_id"] == fin2["bill"]["id"]

def test_generate_invoice_pdf_user_explicitly_mentions_bill_id(test_db):
    from src.agent.context import current_chat_id_var, current_user_message_var
    from src.tools.report_tools import generate_invoice_pdf

    current_chat_id_var.set("chat_test_explicit_text")

    with get_db_connection(test_db) as conn:
        with immediate_transaction(conn):
            add_product_service(conn, {
                "name": "Item Explicit Text", "brand": "BrandET", "unit": "packet", "is_loose": False,
                "hsn_code": "2222", "gst_slab": 5, "cost_price": 50, "sell_price": 100,
                "mrp": 100, "quantity": 10, "reorder_level": 2
            })
            b1 = start_bill_service(conn, "Customer E1", chat_id="chat_test_explicit_text")
            add_item_to_bill_service(conn, b1["bill"]["id"], "Item Explicit Text", quantity=1)
            fin1 = finalize_bill_service(conn, b1["bill"]["id"])

            b2 = start_bill_service(conn, "Customer E2", chat_id="chat_test_explicit_text")
            add_item_to_bill_service(conn, b2["bill"]["id"], "Item Explicit Text", quantity=2)
            fin2 = finalize_bill_service(conn, b2["bill"]["id"])

    b1_id = fin1["bill"]["id"]
    current_user_message_var.set(f"Generate PDF for bill {b1_id}")

    res = generate_invoice_pdf(bill_id=b1_id, db_path=test_db)
    assert res["status"] == "success"
    assert res["bill_id"] == b1_id

def test_generate_invoice_pdf_nonexistent_bill_id_rejected(test_db):
    from src.agent.context import current_chat_id_var, current_user_message_var
    from src.tools.report_tools import generate_invoice_pdf

    current_chat_id_var.set("chat_test_nonexistent")
    current_user_message_var.set("Generate PDF for bill 99")

    with pytest.raises(ValueError, match="Bill #99 not found"):
        generate_invoice_pdf(bill_id=99, db_path=test_db)

def test_generate_invoice_pdf_no_finalized_bill(test_db):
    from src.agent.context import current_chat_id_var, current_user_message_var
    from src.tools.report_tools import generate_invoice_pdf

    current_chat_id_var.set("chat_no_finalized")
    current_user_message_var.set("Generate PDF for my latest bill")

    with pytest.raises(ValueError, match="No finalized bill found for this chat session to generate invoice PDF."):
        generate_invoice_pdf(db_path=test_db)

def test_generate_invoice_pdf_chat_isolation(test_db):
    from src.agent.context import current_chat_id_var, current_user_message_var
    from src.tools.report_tools import generate_invoice_pdf

    with get_db_connection(test_db) as conn:
        with immediate_transaction(conn):
            add_product_service(conn, {
                "name": "Item Iso", "brand": "BrandI", "unit": "packet", "is_loose": False,
                "hsn_code": "3333", "gst_slab": 5, "cost_price": 50, "sell_price": 100,
                "mrp": 100, "quantity": 10, "reorder_level": 2
            })
            b_chat_a = start_bill_service(conn, "Customer A", chat_id="chat_A")
            add_item_to_bill_service(conn, b_chat_a["bill"]["id"], "Item Iso", quantity=1)
            fin_a = finalize_bill_service(conn, b_chat_a["bill"]["id"])

    b_a_id = fin_a["bill"]["id"]
    current_chat_id_var.set("chat_B")
    current_user_message_var.set("Generate PDF for my latest bill")

    with pytest.raises(ValueError, match="No finalized bill found for this chat session to generate invoice PDF."):
        generate_invoice_pdf(db_path=test_db)

    current_user_message_var.set(f"Generate PDF for bill {b_a_id}")
    with pytest.raises(ValueError, match="does not belong to the current chat session"):
        generate_invoice_pdf(bill_id=b_a_id, db_path=test_db)


